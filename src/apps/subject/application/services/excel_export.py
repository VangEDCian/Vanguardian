from collections import OrderedDict
from dataclasses import dataclass
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.subject.infrastructure.repositories.subject_excel_export import (
    DjangoSubjectExcelExportRepository,
)

EXCEL_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
MAX_EXCEL_COLUMNS = 16_384
VISIT_SORT_KEY = "__export_visit_sort_key__"


class SubjectExcelExportSelectionError(ValueError):
    pass


EXPORT_LABEL_LANGUAGE_CODE = "vi"


@dataclass(frozen=True)
class SubjectExcelExportResult:
    content: bytes
    filename: str
    subject_count: int
    field_count: int


class _StudyExportCatalogAdapter:
    @staticmethod
    def list_groups(*, study_id: int) -> list[dict]:
        from apps.study.public import list_subject_export_field_groups

        return list_subject_export_field_groups(
            study_id=study_id,
            language_code=EXPORT_LABEL_LANGUAGE_CODE,
        )


class _DataCaptureExportAdapter:
    @staticmethod
    def read_values(**kwargs) -> dict:
        from apps.datacapture.public import read_subject_export_values

        return read_subject_export_values(**kwargs)


class SubjectExcelExportService:
    repository_class = DjangoSubjectExcelExportRepository
    field_catalog_adapter_class = _StudyExportCatalogAdapter
    data_capture_adapter_class = _DataCaptureExportAdapter

    def __init__(
        self,
        repository=None,
        field_catalog_adapter=None,
        data_capture_adapter=None,
    ):
        self.repository = repository or self.repository_class()
        self.field_catalog_adapter = (
            field_catalog_adapter or self.field_catalog_adapter_class()
        )
        self.data_capture_adapter = (
            data_capture_adapter or self.data_capture_adapter_class()
        )

    def export(
        self,
        *,
        study_id: int,
        site_id: int,
        subject_ids: tuple[int, ...],
        selected_field_tokens: tuple[str, ...],
    ) -> SubjectExcelExportResult:
        selected_fields = self._resolve_selected_fields(
            study_id=study_id,
            selected_field_tokens=selected_field_tokens,
        )
        subject_rows = self.repository.list_scoped_subject_rows(
            study_id=study_id,
            site_id=site_id,
            subject_ids=subject_ids,
        )
        subject_rows = tuple(sorted(subject_rows, key=self._subject_row_sort_key))
        if not subject_rows:
            raise SubjectExcelExportSelectionError(
                "No selected subjects are available in the current study site."
            )

        scoped_subject_ids = tuple(int(row["id"]) for row in subject_rows)
        values_by_subject_id = self.data_capture_adapter.read_values(
            study_id=study_id,
            site_id=site_id,
            subject_ids=scoped_subject_ids,
            field_specs=selected_fields,
        )
        workbook_content = self._build_workbook(
            subject_rows=subject_rows,
            selected_fields=selected_fields,
            values_by_subject_id=values_by_subject_id,
        )
        timestamp = timezone.now().strftime("%Y%m%d-%H%M%S")
        return SubjectExcelExportResult(
            content=workbook_content,
            filename=f"subjects-{study_id}-{timestamp}.xlsx",
            subject_count=len(subject_rows),
            field_count=len(selected_fields),
        )

    def _resolve_selected_fields(
        self,
        *,
        study_id: int,
        selected_field_tokens: tuple[str, ...],
    ) -> tuple[dict, ...]:
        normalized_tokens = tuple(
            dict.fromkeys(
                str(token).strip()
                for token in selected_field_tokens
                if str(token).strip()
            )
        )
        if not normalized_tokens:
            raise SubjectExcelExportSelectionError(
                "Select at least one field to export."
            )

        groups = self.field_catalog_adapter.list_groups(study_id=study_id)
        catalog_fields = tuple(
            field
            for group in groups
            for form in group["forms"]
            for field in form["fields"]
        )
        fields_by_token = {field["token"]: field for field in catalog_fields}
        unknown_tokens = [
            token for token in normalized_tokens if token not in fields_by_token
        ]
        if unknown_tokens:
            raise SubjectExcelExportSelectionError(
                "One or more selected export fields are invalid."
            )
        selected_token_set = set(normalized_tokens)
        selected_fields = tuple(
            field
            for field in catalog_fields
            if field["token"] in selected_token_set
        )
        if len(selected_fields) + 4 > MAX_EXCEL_COLUMNS:
            raise SubjectExcelExportSelectionError(
                "Too many fields were selected for one Excel worksheet."
            )
        return selected_fields

    @classmethod
    def _build_workbook(
        cls,
        *,
        subject_rows: tuple[dict, ...],
        selected_fields: tuple[dict, ...],
        values_by_subject_id: dict[int, dict | tuple[dict, ...]],
    ) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Subjects"
        field_groups = cls._build_field_groups(selected_fields)
        field_columns = tuple(
            column
            for group in field_groups
            for column in group["columns"]
        )
        field_headers = tuple(column["description"] for column in field_columns)
        common_headers = [
            "Subject Code",
            "Screening Code",
            "Randomization Code",
            "Visit",
        ]
        first_header_row = [*common_headers]
        second_header_row = [None] * len(common_headers)
        for group in field_groups:
            column_count = len(group["columns"])
            first_header_row.extend(
                [group["title"], *([None] * (column_count - 1))]
            )
            second_header_row.extend(
                column["description"] for column in group["columns"]
            )

        worksheet.append(first_header_row)
        worksheet.append(second_header_row)
        cls._style_header(worksheet)
        for column_index in range(1, len(common_headers) + 1):
            worksheet.merge_cells(
                start_row=1,
                start_column=column_index,
                end_row=2,
                end_column=column_index,
            )

        group_start_column = len(common_headers) + 1
        for group in field_groups:
            group_end_column = group_start_column + len(group["columns"]) - 1
            if group_end_column > group_start_column:
                worksheet.merge_cells(
                    start_row=1,
                    start_column=group_start_column,
                    end_row=1,
                    end_column=group_end_column,
                )
            group_start_column = group_end_column + 1

        for subject in subject_rows:
            subject_id = int(subject["id"])
            for visit_index, subject_values in enumerate(
                cls._subject_value_rows(
                    values_by_subject_id.get(subject_id),
                ),
                start=1,
            ):
                visit_label = cls._resolve_visit_label(
                    subject_values=subject_values,
                    visit_index=visit_index,
                )
                row_values = [
                    subject.get("subject_code") or "",
                    subject.get("screening_code") or "",
                    subject.get("randomization_code") or "",
                    visit_label,
                    *(
                        cls._field_values_for_column(
                            subject_values=subject_values,
                            field_tokens=column["tokens"],
                        )
                        for column in field_columns
                    ),
                ]
                worksheet.append(row_values)
                cls._force_formula_like_strings_to_text(
                    worksheet,
                    row_index=worksheet.max_row,
                    values=row_values,
                )

        worksheet.freeze_panes = "A3"
        cls._set_column_widths(
            worksheet,
            [*common_headers, *field_headers],
        )
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _resolve_visit_label(*, subject_values: dict, visit_index: int) -> str:
        visit_metadata = subject_values.get(VISIT_SORT_KEY)
        if isinstance(visit_metadata, tuple):
            sequence_no = int(visit_metadata[0] if len(visit_metadata) > 0 else 0)
            event_name = str(
                visit_metadata[5]
                if len(visit_metadata) > 5 and visit_metadata[5] is not None
                else ""
            ).strip()
            form_repeat_index = (
                int(visit_metadata[6])
                if len(visit_metadata) > 6 and visit_metadata[6] is not None
                else None
            )

            label = event_name or (
                "Screening" if sequence_no == 0 else f"Visit {sequence_no}"
            )
            if form_repeat_index is not None:
                return f"{label} #{form_repeat_index}"
            return label

        return f"Visit {visit_index}"

    @staticmethod
    def _subject_row_sort_key(subject_row: dict) -> tuple:
        subject_code = str(subject_row.get("subject_code") or "").strip()
        screening_code = str(subject_row.get("screening_code") or "").strip()
        return (
            0 if subject_code else 1,
            subject_code,
            screening_code,
            int(subject_row.get("id") or 0),
        )

    @staticmethod
    def _subject_value_rows(raw_subject_values) -> tuple[dict, ...]:
        if isinstance(raw_subject_values, dict):
            return (raw_subject_values,)
        if isinstance(raw_subject_values, (list, tuple)):
            rows = [
                row
                for row in raw_subject_values
                if isinstance(row, dict)
            ]
            if rows:
                return tuple(rows)
        return ({},)

    @classmethod
    def _field_description_headers(
        cls,
        selected_fields: tuple[dict, ...],
    ) -> tuple[str, ...]:
        return tuple(
            column["description"]
            for column in cls._build_field_columns(selected_fields)
        )

    @classmethod
    def _build_field_columns(
        cls,
        selected_fields: tuple[dict, ...],
    ) -> tuple[dict, ...]:
        return tuple(
            column
            for group in cls._build_field_groups(selected_fields)
            for column in group["columns"]
        )

    @classmethod
    def _build_field_groups(
        cls,
        selected_fields: tuple[dict, ...],
    ) -> tuple[dict, ...]:
        groups: OrderedDict[tuple, dict] = OrderedDict()
        for field in selected_fields:
            description = cls._field_description(field)
            if not description:
                description = cls._field_fallback_header(field)
            token = str(field.get("token", ""))
            group_key = cls._field_group_key(field, token=token)
            group = groups.setdefault(
                group_key,
                {
                    "title": cls._crf_form_header(field),
                    "columns": [],
                },
            )
            group["columns"].append(
                {
                    "description": description,
                    "tokens": (token,),
                }
            )

        return tuple(groups.values())

    @staticmethod
    def _field_group_key(field: dict, *, token: str) -> tuple:
        binding_id = field.get("binding_id")
        if binding_id is not None:
            return ("binding", binding_id)
        form_identity = (
            field.get("event_definition_id"),
            field.get("crf_template_id"),
            field.get("crf_code"),
        )
        if any(value is not None for value in form_identity):
            return ("form", *form_identity)
        return ("field", token)

    @staticmethod
    def _crf_form_header(field: dict) -> str:
        return str(
            field.get("crf_name")
            or field.get("crf_code")
            or "CRF Form"
        ).strip()

    @staticmethod
    def _field_fallback_header(field: dict) -> str:
        return str(
            field.get("header")
            or field.get("field_key")
            or "Field"
        ).strip()

    @staticmethod
    def _field_values_for_column(
        *,
        subject_values: dict,
        field_tokens: tuple[str, ...],
    ):
        values = []
        for token in field_tokens:
            value = subject_values.get(token)
            if value not in (None, ""):
                values.append(value)

        if not values:
            return ""
        if len(values) == 1:
            return values[0]
        return "; ".join(str(value) for value in values if str(value).strip())

    @staticmethod
    def _field_description(field: dict) -> str:
        return str(
            field.get("field_label")
            or field.get("header")
            or field.get("field_key")
            or "Field"
        ).strip()

    @staticmethod
    def _style_header(worksheet):
        fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        for row in worksheet.iter_rows(min_row=1, max_row=2):
            for cell in row:
                cell.font = Font(bold=True, color="17324D")
                cell.fill = fill
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

    @staticmethod
    def _force_formula_like_strings_to_text(worksheet, *, row_index, values):
        for column_index, value in enumerate(values, start=1):
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                ).data_type = "s"

    @staticmethod
    def _set_column_widths(worksheet, headers):
        for column_index, header in enumerate(headers, start=1):
            worksheet.column_dimensions[
                get_column_letter(column_index)
            ].width = min(max(len(str(header)) + 2, 14), 60)


__all__ = [
    "EXCEL_CONTENT_TYPE",
    "SubjectExcelExportResult",
    "SubjectExcelExportSelectionError",
    "SubjectExcelExportService",
]
