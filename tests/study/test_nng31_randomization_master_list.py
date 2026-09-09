import csv
import io
from collections import Counter
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management.base import CommandError
from django.test import SimpleTestCase, TestCase
from django.utils import timezone
from openpyxl import load_workbook

from apps.study.application.commands.import_randomization.types import (
    PreviewRandomizationImportCommand,
    RandomizationImportValidationError,
)
from apps.study.application.commands.nng31_master_list import (
    ApproveNng31MasterListCommand,
    CommitNng31MasterListCommand,
)
from apps.study.application.services.import_randomization_commit import (
    _ensure_master_list_unlocked,
)
from apps.study.application.services.nng31_master_list import (
    ApproveNng31MasterListService,
    CommitNng31MasterListImportService,
    PreviewNng31MasterListImportService,
)
from apps.study.domain import format_randomization_code
from apps.study.infrastructure.repositories.nng31_master_list import (
    DjangoNng31MasterListRepository,
)
from apps.study.management.randomization_master_list import (
    REQUIRED_COLUMNS,
    parse_and_validate_nng31_master_list,
)
from apps.study.models import (
    RandomizationArm,
    RandomizationEvent,
    RandomizationScheme,
    RandomizationSlot,
    Study,
)
from apps.study.presentation.web.forms import Nng31MasterListImportFileForm


class Nng31RandomizationMasterListTests(SimpleTestCase):
    STATISTICIAN_ARM_SEQUENCE = (
        "NEEN"
        "NENN"
        "EEEN"
        "EENN"
        "ENNE"
        "EENNNN"
        "EEEENE"
        "NNNENE"
        "ENENEN"
    )

    def test_randomization_code_uses_the_configured_prefix_literally(self):
        self.assertEqual(
            format_randomization_code(prefix="NNG31-", sequence_no=1, padding=3),
            "NNG31-001",
        )
        self.assertEqual(
            format_randomization_code(prefix="NNG31-R", sequence_no=1, padding=3),
            "NNG31-R001",
        )

    def test_static_template_is_an_excel_workbook_with_44_allocations(self):
        workbook = load_workbook(self._template_path(), read_only=True, data_only=True)
        rows = list(workbook.worksheets[0].iter_rows(values_only=True))

        self.assertEqual(tuple(rows[0]), REQUIRED_COLUMNS)
        self.assertEqual(len(rows), 45)
        self.assertEqual(rows[1][0], "CROSS_OVER_NANOKINE_EPREX4000IU")
        self.assertEqual(rows[1][1], "NNG31-001")
        self.assertEqual(rows[1][4], "SEQ_NANOKINE_EPREX4000IU")
        self.assertEqual(rows[-1][1], "NNG31-044")

    def test_end_user_form_accepts_xlsx_and_rejects_csv(self):
        xlsx_form = Nng31MasterListImportFileForm(
            data={"master_list_version": "1.0"},
            files={
                "import_file": SimpleUploadedFile(
                    "master-list.xlsx",
                    self._template_path().read_bytes(),
                )
            },
        )
        csv_form = Nng31MasterListImportFileForm(
            data={"master_list_version": "1.0"},
            files={"import_file": SimpleUploadedFile("master-list.csv", self._valid_csv())},
        )

        self.assertTrue(xlsx_form.is_valid())
        self.assertFalse(csv_form.is_valid())

    def test_generic_import_cannot_modify_a_locked_master_list_scheme(self):
        with self.assertRaises(RandomizationImportValidationError):
            _ensure_master_list_unlocked(
                scheme=SimpleNamespace(master_list_locked_at=object()),
                parsed_row=SimpleNamespace(row_number=2, identifier="NNG31_XOVER"),
            )

    def test_accepts_balanced_master_list(self):
        rows, checksum = parse_and_validate_nng31_master_list(
            content=self._valid_csv(),
            study=SimpleNamespace(code="NNG31"),
            scheme=self._scheme(),
            arms_by_code={"SEQ_E_N": object(), "SEQ_N_E": object()},
        )

        self.assertEqual(len(rows), 44)
        self.assertEqual(rows[0].randomization_code, "NNG31-001")
        self.assertEqual(rows[-1].randomization_code, "NNG31-044")
        self.assertEqual(len(checksum), 64)

    def test_accepts_statistician_sequence_without_enforcing_block_shape(self):
        rows, checksum = parse_and_validate_nng31_master_list(
            content=self._statistician_csv(),
            study=SimpleNamespace(code="NNG31"),
            scheme=self._scheme(),
            arms_by_code={"SEQ_E_N": object(), "SEQ_N_E": object()},
        )

        self.assertEqual(len(rows), 44)
        self.assertEqual(Counter(row.arm_code for row in rows)["SEQ_E_N"], 22)
        self.assertEqual(len(checksum), 64)

    def test_accepts_configured_scheme_and_arm_codes_without_name_assumptions(self):
        content = (
            self._statistician_csv()
            .decode("utf-8")
            .replace("NNG31_XOVER", "CROSS_OVER_NANOKINE_EPREX4000IU")
            .replace("SEQ_E_N", "SEQ_EPREX4000IU_NANOKINE")
            .replace("SEQ_N_E", "SEQ_NANOKINE_EPREX4000IU")
            .encode("utf-8")
        )
        scheme = self._scheme()
        scheme.code = "CROSS_OVER_NANOKINE_EPREX4000IU"

        rows, _checksum = parse_and_validate_nng31_master_list(
            content=content,
            study=SimpleNamespace(code="NNG31"),
            scheme=scheme,
            arms_by_code={
                "SEQ_EPREX4000IU_NANOKINE": object(),
                "SEQ_NANOKINE_EPREX4000IU": object(),
            },
        )

        self.assertEqual(len(rows), 44)

    def test_rejects_unequal_overall_sequence_allocation(self):
        content = self._valid_csv().decode("utf-8").replace(
            "NNG31_XOVER,NNG31-002,2,1,SEQ_N_E",
            "NNG31_XOVER,NNG31-002,2,1,SEQ_E_N",
        )

        with self.assertRaisesMessage(
            CommandError,
            "allocate subjects equally between the two sequences",
        ):
            parse_and_validate_nng31_master_list(
                content=content.encode("utf-8"),
                study=SimpleNamespace(code="NNG31"),
                scheme=self._scheme(),
                arms_by_code={"SEQ_E_N": object(), "SEQ_N_E": object()},
            )

    def test_rejects_legacy_randomization_identifier(self):
        content = self._valid_csv().decode("utf-8").replace("NNG31-001", "R-001")

        with self.assertRaisesMessage(CommandError, "NNG31-001 through NNG31-044"):
            parse_and_validate_nng31_master_list(
                content=content.encode("utf-8"),
                study=SimpleNamespace(code="NNG31"),
                scheme=self._scheme(),
                arms_by_code={"SEQ_E_N": object(), "SEQ_N_E": object()},
            )

    @staticmethod
    def _valid_csv():
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=REQUIRED_COLUMNS, lineterminator="\n")
        writer.writeheader()
        sequence_no = 0
        for block_no, block_size in enumerate((4, 4, 4, 4, 4, 6, 6, 6, 6), start=1):
            for offset in range(block_size):
                sequence_no += 1
                writer.writerow(
                    {
                        "Scheme Code": "NNG31_XOVER",
                        "Randomization ID": f"NNG31-{sequence_no:03}",
                        "Sequence No": sequence_no,
                        "Block No": block_no,
                        "Arm Code": "SEQ_E_N" if offset % 2 == 0 else "SEQ_N_E",
                    }
                )
        return output.getvalue().encode("utf-8")

    @classmethod
    def _statistician_csv(cls):
        source = io.StringIO(cls._valid_csv().decode("utf-8"))
        rows = list(csv.DictReader(source))
        arm_codes = {"E": "SEQ_E_N", "N": "SEQ_N_E"}
        for row, sequence_arm in zip(
            rows,
            cls.STATISTICIAN_ARM_SEQUENCE,
            strict=True,
        ):
            row["Arm Code"] = arm_codes[sequence_arm]

        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=REQUIRED_COLUMNS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
        return output.getvalue().encode("utf-8")

    @staticmethod
    def _scheme():
        return SimpleNamespace(
            code="NNG31_XOVER",
            target_randomized_total=44,
            randomization_code_prefix="NNG31-",
            randomization_code_padding=3,
        )

    @staticmethod
    def _template_path():
        return (
            Path(__file__).resolve().parents[2]
            / "src/staticfiles/study/templates/nng31_randomization_master_list_template.xlsx"
        )

    @classmethod
    def _statistician_template_content(cls):
        workbook = load_workbook(cls._template_path())
        worksheet = workbook.worksheets[0]
        arm_codes = {
            "E": "SEQ_EPREX4000IU_NANOKINE",
            "N": "SEQ_NANOKINE_EPREX4000IU",
        }
        for row_number, sequence_arm in enumerate(
            cls.STATISTICIAN_ARM_SEQUENCE,
            start=2,
        ):
            worksheet.cell(row=row_number, column=5, value=arm_codes[sequence_arm])
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()


class Nng31MasterListEndUserImportServiceTests(TestCase):
    def setUp(self):
        self.study = SimpleNamespace(pk=3, code="NNG31")
        self.scheme = SimpleNamespace(
            pk=7,
            code="CROSS_OVER_NANOKINE_EPREX4000IU",
            target_randomized_total=44,
            randomization_code_prefix="NNG31-",
            randomization_code_padding=3,
            master_list_locked_at=None,
        )
        self.arm_e_n = SimpleNamespace(
            pk=11,
            arm_code="SEQ_EPREX4000IU_NANOKINE",
        )
        self.arm_n_e = SimpleNamespace(
            pk=12,
            arm_code="SEQ_NANOKINE_EPREX4000IU",
        )
        self.periods = (
            self._period(self.arm_e_n, 1, "EPREX_4000U"),
            self._period(self.arm_e_n, 2, "NANOKINE"),
            self._period(self.arm_n_e, 1, "NANOKINE"),
            self._period(self.arm_n_e, 2, "EPREX_4000U"),
        )

    def test_preview_resolves_treatment_sequence_from_configured_periods(self):
        repository = self._repository()
        template_path = Nng31RandomizationMasterListTests._template_path()
        result = PreviewNng31MasterListImportService(repository=repository).execute(
            PreviewRandomizationImportCommand(
                actor_user_id=9,
                study_id=self.study.pk,
                file_name=template_path.name,
                file_content=template_path.read_bytes(),
            )
        )

        self.assertEqual(result.issues, ())
        self.assertEqual(result.total_rows, 44)
        self.assertEqual(result.columns[-1].label, "Treatment Sequence")
        self.assertEqual(result.preview_rows[0].values[-1], "NANOKINE → EPREX_4000U")
        self.assertEqual(result.preview_rows[2].values[-1], "EPREX_4000U → NANOKINE")
        self.assertEqual(len(result.checksum), 64)

    def test_preview_accepts_statistician_sequence_without_block_shape_rules(self):
        result = PreviewNng31MasterListImportService(
            repository=self._repository()
        ).execute(
            PreviewRandomizationImportCommand(
                actor_user_id=99,
                study_id=self.study.pk,
                file_name="nng31_randomization_master_list_template.xlsx",
                file_content=(
                    Nng31RandomizationMasterListTests._statistician_template_content()
                ),
            )
        )

        self.assertEqual(result.issues, ())
        self.assertEqual(result.total_rows, 44)

    def test_missing_code_prefix_produces_one_configuration_issue(self):
        self.scheme.randomization_code_prefix = ""
        repository = self._repository()
        template_path = Nng31RandomizationMasterListTests._template_path()

        result = PreviewNng31MasterListImportService(repository=repository).execute(
            PreviewRandomizationImportCommand(
                actor_user_id=9,
                study_id=self.study.pk,
                file_name=template_path.name,
                file_content=template_path.read_bytes(),
            )
        )

        prefix_issues = [
            issue
            for issue in result.issues
            if issue.column_label == "Randomization ID"
            and "Prefix" in str(issue.reason)
        ]
        self.assertEqual(len(prefix_issues), 1)

    def test_commit_imports_slots_and_records_version_and_checksum(self):
        repository = self._repository()
        template_path = Nng31RandomizationMasterListTests._template_path()
        repository.list_slots.return_value = []
        repository.update_or_create_slot.side_effect = [
            (SimpleNamespace(), True) for _index in range(44)
        ]
        audit_adapter = MagicMock()
        preview_service = PreviewNng31MasterListImportService(repository=repository)
        service = CommitNng31MasterListImportService(
            preview_service=preview_service,
            repository=repository,
            audit_adapter=audit_adapter,
        )

        result = service.execute(
            CommitNng31MasterListCommand(
                actor_user_id=9,
                study_id=self.study.pk,
                file_name=template_path.name,
                file_content=template_path.read_bytes(),
                master_list_version="1.0",
            )
        )

        self.assertEqual(result.created_count, 44)
        self.assertEqual(repository.update_or_create_slot.call_count, 44)
        self.assertEqual(self.scheme.master_list_version, "1.0")
        self.assertEqual(self.scheme.master_list_checksum, result.checksum)
        self.assertEqual(self.scheme.master_list_source_filename, template_path.name)
        repository.save_scheme.assert_called_once()
        audit_adapter.record_event.assert_called_once()

    def test_commit_bypasses_assigned_arm_and_leaves_18_22_available(self):
        repository = self._repository()
        template_path = Nng31RandomizationMasterListTests._template_path()
        existing_slots = []
        for sequence_no in range(1, 45):
            arm = self.arm_n_e if sequence_no <= 22 else self.arm_e_n
            existing_slots.append(
                SimpleNamespace(
                    pk=100 + sequence_no,
                    sequence_no=sequence_no,
                    status="assigned" if sequence_no <= 4 else "available",
                    arm_id=arm.pk,
                    arm=arm,
                    randomization_code=None,
                )
            )
        existing_slots.extend(
            [
                SimpleNamespace(
                    pk=200 + sequence_no,
                    sequence_no=sequence_no,
                    status="void",
                    arm_id=self.arm_e_n.pk,
                    arm=self.arm_e_n,
                    randomization_code=None,
                )
                for sequence_no in (45, 46)
            ]
        )
        repository.list_slots.return_value = existing_slots

        def replace_existing_slots(*, scheme, slot_rows, retired_slots, now):
            self.assertEqual(scheme, self.scheme)
            self.assertEqual(len(retired_slots), 2)
            self.assertTrue(all(slot.status == "void" for slot in retired_slots))
            assigned_rows = [
                (slot, values)
                for slot, values in slot_rows
                if slot.status == "assigned"
            ]
            self.assertEqual(len(slot_rows), 44)
            self.assertEqual(len(assigned_rows), 4)
            self.assertTrue(
                all(values["arm"].pk == self.arm_n_e.pk for _slot, values in assigned_rows)
            )
            available_rows = [
                values
                for slot, values in slot_rows
                if slot.status != "assigned"
            ]
            self.assertEqual(
                sum(values["arm"].pk == self.arm_n_e.pk for values in available_rows),
                18,
            )
            self.assertEqual(
                sum(values["arm"].pk == self.arm_e_n.pk for values in available_rows),
                22,
            )
            return tuple(
                {
                    "slot_id": slot.pk,
                    "arm_id": values["arm"].pk,
                    "arm_code": values["arm"].arm_code,
                    "randomization_number": values["randomization_code"],
                }
                for slot, values in assigned_rows
            )

        repository.replace_existing_slots.side_effect = replace_existing_slots
        subject_slot_reconciler = MagicMock()
        service = CommitNng31MasterListImportService(
            preview_service=PreviewNng31MasterListImportService(repository=repository),
            repository=repository,
            audit_adapter=MagicMock(),
            subject_slot_reconciler=subject_slot_reconciler,
        )

        result = service.execute(
            CommitNng31MasterListCommand(
                actor_user_id=9,
                study_id=self.study.pk,
                file_name=template_path.name,
                file_content=template_path.read_bytes(),
                master_list_version="1.0",
            )
        )

        self.assertEqual(result.updated_count, 44)
        self.assertEqual(result.created_count, 0)
        self.assertEqual(result.assigned_bypass_count, 4)
        repository.update_or_create_slot.assert_not_called()
        subject_slot_reconciler.assert_called_once()
        self.assertEqual(
            len(subject_slot_reconciler.call_args.kwargs["assignments"]),
            4,
        )

    def test_approval_requires_a_different_user_and_locks_valid_master_list(self):
        repository = self._repository()
        self.scheme.master_list_checksum = "a" * 64
        self.scheme.master_list_version = "1.0"
        self.scheme.master_list_imported_by_id = 9
        repository.list_slots.return_value = self._slots()
        service = ApproveNng31MasterListService(
            preview_service=PreviewNng31MasterListImportService(repository=repository),
            repository=repository,
            audit_adapter=MagicMock(),
        )

        with self.assertRaises(RandomizationImportValidationError):
            service.execute(
                ApproveNng31MasterListCommand(
                    actor_user_id=9,
                    study_id=self.study.pk,
                    scheme_id=self.scheme.pk,
                    expected_checksum=self.scheme.master_list_checksum,
                )
            )

        result = service.execute(
            ApproveNng31MasterListCommand(
                actor_user_id=10,
                study_id=self.study.pk,
                scheme_id=self.scheme.pk,
                expected_checksum=self.scheme.master_list_checksum,
            )
        )

        self.assertEqual(result.scheme_id, self.scheme.pk)
        self.assertEqual(self.scheme.master_list_approved_by_id, 10)
        self.assertIsNotNone(self.scheme.master_list_locked_at)

    def _repository(self):
        repository = MagicMock()
        repository.get_study.return_value = self.study
        repository.get_scheme.return_value = self.scheme
        repository.list_active_arms.return_value = [self.arm_e_n, self.arm_n_e]
        repository.list_sequence_periods.return_value = self.periods
        return repository

    def _slots(self):
        template_path = Nng31RandomizationMasterListTests._template_path()
        preview = PreviewNng31MasterListImportService(
            repository=self._repository()
        ).execute(
            PreviewRandomizationImportCommand(
                actor_user_id=9,
                study_id=self.study.pk,
                file_name=template_path.name,
                file_content=template_path.read_bytes(),
            )
        )
        arm_by_code = {
            self.arm_e_n.arm_code: self.arm_e_n,
            self.arm_n_e.arm_code: self.arm_n_e,
        }
        return [
            SimpleNamespace(
                randomization_code=row.values["randomization_code"],
                sequence_no=row.values["sequence_no"],
                block_no=row.values["block_no"],
                arm=arm_by_code[row.values["arm_code"]],
            )
            for row in preview.parsed_rows
        ]

    @staticmethod
    def _period(arm, period_no, treatment_code):
        return SimpleNamespace(
            arm=arm,
            period_no=period_no,
            treatment_code=treatment_code,
            start_event_definition_id=100 + period_no,
            end_event_definition_id=200 + period_no,
        )


class Nng31MasterListRepositoryTests(TestCase):
    def test_replace_existing_slots_avoids_unique_conflicts_and_preserves_assignment(self):
        now = timezone.now()
        study = Study.objects.create(
            created_at=now,
            updated_at=now,
            code="NNG31-REMAP",
            name="NNG31 remap test",
        )
        scheme = RandomizationScheme.objects.create(
            created_at=now,
            updated_at=now,
            study=study,
            code="CROSS_OVER_NANOKINE_EPREX4000IU",
            name="Crossover",
            randomization_type="stratified_blocked",
            target_randomized_total=2,
        )
        arm_n_e = RandomizationArm.objects.create(
            created_at=now,
            updated_at=now,
            scheme=scheme,
            arm_code="SEQ_NANOKINE_EPREX4000IU",
            arm_name="NANOKINE then EPREX",
            target_count=1,
        )
        arm_e_n = RandomizationArm.objects.create(
            created_at=now,
            updated_at=now,
            scheme=scheme,
            arm_code="SEQ_EPREX4000IU_NANOKINE",
            arm_name="EPREX then NANOKINE",
            target_count=1,
        )
        assigned_slot = RandomizationSlot.objects.create(
            created_at=now,
            updated_at=now,
            scheme=scheme,
            arm=arm_n_e,
            sequence_no=1,
            randomization_code="R-001",
            status="assigned",
            assigned_subject_id=501,
            assigned_event_id=601,
            assigned_at=now,
        )
        available_slot = RandomizationSlot.objects.create(
            created_at=now,
            updated_at=now,
            scheme=scheme,
            arm=arm_e_n,
            sequence_no=2,
            randomization_code="R-002",
            status="void",
            assigned_subject_id=999,
        )
        surplus_slot = RandomizationSlot.objects.create(
            created_at=now,
            updated_at=now,
            scheme=scheme,
            arm=arm_e_n,
            sequence_no=3,
            randomization_code="R-003",
            status="void",
        )
        event = RandomizationEvent.objects.create(
            created_at=now,
            event_type="Assigned",
            randomization_source="workflow_action",
            randomization_sequence=arm_n_e.arm_code,
            randomization_number="R-001",
            subject_id=501,
            study=study,
            scheme=scheme,
            arm=arm_n_e,
            slot=assigned_slot,
        )

        metadata = DjangoNng31MasterListRepository().replace_existing_slots(
            scheme=scheme,
            slot_rows=[
                (
                    assigned_slot,
                    {
                        "sequence_no": 2,
                        "arm": arm_n_e,
                        "randomization_code": "NNG31-002",
                        "block_no": 1,
                    },
                ),
                (
                    available_slot,
                    {
                        "sequence_no": 1,
                        "arm": arm_e_n,
                        "randomization_code": "NNG31-001",
                        "block_no": 1,
                    },
                ),
            ],
            retired_slots=[surplus_slot],
            now=now,
        )

        assigned_slot.refresh_from_db()
        available_slot.refresh_from_db()
        surplus_slot.refresh_from_db()
        event.refresh_from_db()
        self.assertEqual(assigned_slot.sequence_no, 2)
        self.assertEqual(assigned_slot.randomization_code, "NNG31-002")
        self.assertEqual(assigned_slot.assigned_subject_id, 501)
        self.assertEqual(assigned_slot.status, "assigned")
        self.assertEqual(available_slot.sequence_no, 1)
        self.assertEqual(available_slot.randomization_code, "NNG31-001")
        self.assertEqual(available_slot.status, "available")
        self.assertIsNone(available_slot.assigned_subject_id)
        self.assertTrue(surplus_slot.deleted)
        self.assertIsNone(surplus_slot.randomization_code)
        self.assertGreater(surplus_slot.sequence_no, 3)
        self.assertEqual(event.randomization_number, "NNG31-002")
        self.assertEqual(metadata[0]["slot_id"], assigned_slot.pk)
